import re, json, html, urllib.request
from concurrent.futures import ThreadPoolExecutor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
B="http://inheritance.ilmsummit.org/projects/inheritance/"
def get(u):
    for _ in range(3):
        try: return urllib.request.urlopen(u,timeout=150).read().decode('utf8','replace')
        except Exception as e: err=e
    raise err
def txt(s): return re.sub(r'\s+',' ',html.unescape(re.sub(r'<[^>]+>',' ',s))).strip()
lst=get(B+"testcasespage.aspx")
cases=[]
ids=sorted({int(x) for x in re.findall(r'TestCaseID=(\d+)',lst)})
# description: spans following each link
desc={}
for m in re.finditer(r'TestCaseID=(\d+)".*?<span[^>]*>\s*\((.*?)\)\s*</span>',lst,re.S):
    desc.setdefault(int(m.group(1)),txt(m.group(2)))
def table(h,tid):
    i=h.find('id="%s"'%tid)
    if i<0: return []
    j=h.find('</table>',i)
    rows=[]
    for tr in re.findall(r'<tr.*?</tr>',h[i:j],re.S)[1:]:
        c=[txt(x) for x in re.findall(r'<td.*?</td>',tr,re.S)]
        rows.append(c)
    return rows
def one(i):
    h=get(B+"Results.aspx?TestCaseID=%d"%i)
    steps=[txt(x) for x in re.findall(r'<li>(.*?)</li>',h,re.S)]
    return dict(id=i,heirs=desc.get(i,''),category=table(h,'dgSharesCtg'),individual=table(h,'dgSharesIndiv'),steps=steps)
with ThreadPoolExecutor(16) as ex: data=list(ex.map(one,ids))
json.dump(data,open('data/ilmsummit_cases.json','w',encoding='utf8'),ensure_ascii=False,indent=1)
wb=Workbook(); ws=wb.active; ws.title="Test Cases"
ws.append(["Case #","Heirs (input)","Category","Fraction","Percent","Calculation Steps (reference)"])
hf=PatternFill("solid",fgColor="1C5E55")
for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=hf
w2=wb.create_sheet("By Individual"); w2.append(["Case #","Heirs (input)","Relative","Fraction","Percent"])
for c in w2[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=hf
for d in data:
    st="\n".join("%d. %s"%(k+1,s) for k,s in enumerate(d['steps']))
    for k,r in enumerate(d['category']):
        ws.append([d['id'],d['heirs'],*r[:3],st if k==0 else ""])
    for r in d['individual']: w2.append([d['id'],d['heirs'],*r[:3]])
for w,wd in ((ws,[8,45,22,10,10,110]),(w2,[8,45,25,10,10])):
    for n,x in enumerate(wd): w.column_dimensions[chr(65+n)].width=x
    for row in w.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
wb.save('Inheritance_TestCases_ilmsummit.xlsx')
print(len(data),sum(1 for d in data if not d['individual']),"empty")


